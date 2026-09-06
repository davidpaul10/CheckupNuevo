import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = docx.Document()

# Page Setup
for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

# Title
title_p = doc.add_paragraph()
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_title = title_p.add_run('GUÍA COMPLETA DE PRESENTACIÓN Y DEFENSA DEL PROYECTO: CHECKUP+\n')
run_title.bold = True
run_title.font.size = Pt(18)
run_title.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)

sub_p = doc.add_paragraph()
sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_sub = sub_p.add_run('Evaluación de Cierre de la Semana 4 (QA y Despliegue de MVP)\nEcosistema Tecnológico - Proyecto CheckUp+\n')
run_sub.font.size = Pt(11)
run_sub.font.italic = True
run_sub.font.color.rgb = RGBColor(0x59, 0x59, 0x59)

doc.add_paragraph().paragraph_format.space_after = Pt(12)

def add_heading_1(text):
    h = doc.add_paragraph()
    r = h.add_run(text)
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(6)
    return h

def add_heading_2(text):
    h = doc.add_paragraph()
    r = h.add_run(text)
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(0x0E, 0x7C, 0x7B)
    h.paragraph_format.space_before = Pt(10)
    h.paragraph_format.space_after = Pt(4)
    return h

def add_body(text, bold=False, italic=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(0x26, 0x26, 0x26)
    p.paragraph_format.space_after = Pt(6)
    return p

def add_bullet(bold_prefix, text):
    p = doc.add_paragraph(style='List Bullet')
    r1 = p.add_run(bold_prefix)
    r1.bold = True
    r1.font.size = Pt(11)
    r2 = p.add_run(text)
    r2.font.size = Pt(11)
    p.paragraph_format.space_after = Pt(4)
    return p

def add_script_box(title, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.right_indent = Inches(0.3)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(8)
    
    r_title = p.add_run(f'🗣️ {title}\n')
    r_title.bold = True
    r_title.font.size = Pt(11)
    r_title.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
    
    r_text = p.add_run(f'"{text}"')
    r_text.italic = True
    r_text.font.size = Pt(10.5)
    r_text.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

# 1. RESUMEN DEL CONTEXTO
add_heading_1('1. CONTEXTO DE LA EVALUACIÓN (SEMANA 4)')
add_body('En cumplimiento con el Manual Técnico General del ecosistema y la Especificación de Requerimientos (spec.md), la evaluación del sábado corresponde al cierre de la Semana 4: QA y Despliegue de MVP.')
add_bullet('Objetivo Principal: ', 'Demostrar que el equipo cuenta con un Producto Mínimo Viable (MVP) completamente funcional, explicar la profundidad de los prompts utilizados, los aprendizajes adquiridos con la Inteligencia Artificial y confirmar que la aplicación está lista para desplegarse en los servidores de Linkiar.')
add_bullet('Evolución del Equipo: ', 'Demostrar cómo se pasó del rol de "programador tradicional" al rol de "Director de IA" (Nivel 3/4 de la Pista de Dirección de IA), utilizando el stack oficial (NotebookLM -> Spec Kit -> Google Stitch -> Antigravity).')

# 2. GUION DE PRESENTACIÓN PASO A PASO
add_heading_1('2. GUION OFICIAL DE EXPOSICIÓN (PASO A PASO)')
add_body('Sugerencia: Dividir estos 4 bloques entre los integrantes del grupo para una presentación fluida de 10 a 12 minutos.')

add_heading_2('Bloque 1: Introducción y Recorrido por las Semanas (1.5 minutos)')
add_script_box('Guion para el Integrante 1', 
'Buenos días/tardes a todos. Presentamos el estado del proyecto CheckUp+ al cierre de la Semana 4 (QA y Despliegue de MVP).\n\n'
'Siguiendo la metodología oficial de 8 semanas del manual, nuestro avance fue el siguiente:\n'
'• Semana 1 (Investigación en tiempo real): Consolidamos en NotebookLM entrevistas de campo y análisis competitivo de laboratorios en Quito. Identificamos los 3 problemas centrales: ineficiencia logística en recepción, estigma e incomodidad en salud sexual y terminología médica incomprensible.\n'
'• Semana 2 (Especificación): Tradujimos la investigación en nuestro spec.md usando Spec Kit y diseñamos las pantallas clave en Google Stitch.\n'
'• Semana 3 (Desarrollo Core): Dirigimos agentes de desarrollo en Antigravity para construir la arquitectura en Angular 22.\n'
'• Semana 4 (QA y MVP): Ejecutamos el control de calidad de flujos clave y dejamos la aplicación lista para su despliegue en los servidores de Linkiar.')

add_heading_2('Bloque 2: Demostración en Vivo del MVP (4 minutos)')
add_script_box('Guion para el Integrante 2 (con la app en pantalla)', 
'En nuestra aplicación CheckUp+ hemos implementado los requerimientos del spec.md:\n\n'
'1. Módulo 4 - Pase Verde de Salud Sexual (RF4.1 y RF4.2):\n'
'Aquí ven el Pase Verde con código QR dinámico y un temporizador en tiempo real de 15 minutos. Cumplimos estrictamente el RF4.2 de Privacidad Visual: si un tercero escanea el QR, solo verá el estado "Al Día" y la fecha de vigencia; NUNCA se muestran diagnósticos, resultados ni nombres de enfermedades.\n'
'Además, cuenta con cámara nativa para escaneo y descarga del certificado médico oficial en PDF.\n\n'
'2. Módulo 2 - Reserva Express en menos de 3 pasos (RF2.2 y RF2.3):\n'
'El usuario busca laboratorios cercanos (San José, Clínica de la Mujer), selecciona el examen y confirma la cita en menos de 3 clics, generando un QR de Cita (Fast-Track) para ingresar al laboratorio sin filas ni papeleo.\n\n'
'3. Módulos 1 y 3 - Perfil y Seguridad (RF1.2 y RF3.2):\n'
'Incluye historial centralizado y botón de Modo Confidencialidad para proteger datos sensibles.')

add_heading_2('Bloque 3: Profundidad Técnica y Prompting con IA (3 minutos)')
add_script_box('Guion para el Integrante 3', 
'Para lograr este MVP, actuamos como Directores de IA (Nivel 3/4 del manual):\n\n'
'• Profundidad de Prompts: No usamos prompts genéricos de una sola línea. Redactamos prompts técnicos estructurados asignando reglas de negocio, tipado estricto en TypeScript y restricciones UI/UX.\n'
'• Bitácora prompts.md: Documentamos los prompts óptimos en la raíz del repositorio para que otros equipos del ecosistema reutilicen nuestras soluciones.\n'
'• Arquitectura Angular 22: Implementamos Angular Signals (signal()) para manejo de estado reactivo ultra-rápido y consumo directo de Web APIs nativas (cámara y AudioContext).')

add_heading_2('Bloque 4: QA, RNF y Despliegue en Linkiar (1.5 minutos)')
add_script_box('Guion para el Integrante 4', 
'Para cerrar la Semana 4:\n'
'• Validamos los Requerimientos No Funcionales: generación y validación de QR en menos de 2 segundos (RNF3) y cifrado de datos (RNF1).\n'
'• Etiquetamos la versión en el repositorio como sprint16-semana4.\n'
'• El código está empaquetado, verificado y listo para ser subido a los servidores de Linkiar para dar paso a la etapa comercial de las Semanas 5 y 6.')

# 3. CONCEPTOS CLAVE
add_heading_1('3. CONCEPTOS TÉCNICOS QUE DEBEN APRENDERSE DE MEMORIA')
add_bullet('1. Angular Signals (signal()): ', 'Mecanismo reactivo moderno de Angular 22 que actualiza la pantalla al instante al cambiar un dato (como el temporizador del Pase Verde), sin sobrecargar el procesador.')
add_bullet('2. RF4.2 (Privacidad Visual): ', 'Regla fundamental del Pase Verde. Al escanear el QR, ÚNICAMENTE muestra "Al Día" y vigencia. Jamás expone nombres de enfermedades, pruebas ni diagnósticos.')
add_bullet('3. Check-In Express (QR de Cita): ', 'Flujo del RF2.3 que elimina el registro manual y las filas en recepción al llegar al laboratorio.')
add_bullet('4. Director de IA vs. Programador Tradicional: ', 'Cambio de mentalidad del manual. En lugar de memorizar sintaxis a mano, diriges la construcción mediante especificación (Spec Kit) y prompts estructurados en Antigravity.')
add_bullet('5. Bitácora prompts.md: ', 'Registro obligatorio en el repositorio donde el equipo guarda los prompts más efectivos que resolvieron tareas complejas.')
add_bullet('6. Staging / Producción en Linkiar: ', 'Estado en el que el código ya superó el QA y está empaquetado con sus variables de entorno listo para subir a los servidores de Linkiar.')

# 4. PREGUNTAS Y RESPUESTAS
add_heading_1('4. PREGUNTAS TRAMPA O DIFÍCILES Y SUS RESPUESTAS EXACTAS')

add_heading_2('Pregunta 1: "¿Cómo fue su proceso de trabajo con la Inteligencia Artificial desde la Semana 1?"')
add_body('Respuesta Exacta: "No empezamos programando de inmediato. En la Semana 1 subimos audios de entrevistas y datos de mercado a NotebookLM. En la Semana 2 usamos Spec Kit para traducir esos hallazgos en nuestro spec.md y diseñamos las pantallas en Google Stitch. En las Semanas 3 y 4 usamos Antigravity para dirigir agentes de IA que generaron el código en Angular 22 y ejecutaron las pruebas de QA."', bold=True)

add_heading_2('Pregunta 2: "¿Qué tan profundos o complejos fueron los prompts que utilizaron?"')
add_body('Respuesta Exacta: "Fueron prompts de arquitectura e integración. Por ejemplo, para el Pase Verde especificamos un componente standalone en Angular 22, utilizando Signals para el contador en tiempo real de 15 minutos, integrando getUserMedia para la cámara, AudioContext para el sonido de validación y aplicando los principios de privacidad visual del RF4.2."', bold=True)

add_heading_2('Pregunta 3: "¿Cómo garantizan que el Pase Verde no viole la privacidad del paciente?"')
add_body('Respuesta Exacta: "Por dos mecanismos: El QR es dinámico y expira en 15 minutos (RF4.1), por lo que una foto anterior no sirve. Además, por el RF4.2, al escanearlo el sistema solo responde con el estado \'Al Día\' y la fecha de vigencia. Jamás muestra nombres de exámenes ni diagnósticos."', bold=True)

add_heading_2('Pregunta 4: "¿El código está listo para subirse a los servidores de Linkiar?"')
add_body('Respuesta Exacta: "Sí. Al estar en la Semana 4 (QA y Despliegue), hemos aislado las variables de entorno, el proyecto compila sin errores (ng build), cumple con la estructura standalone y el etiquetado en Git (sprint16-semana4), por lo que está listo para el despliegue en Linkiar."', bold=True)

add_heading_2('Pregunta 5: "¿Qué aprendieron de nuevo utilizando la IA en este proyecto?"')
add_body('Respuesta Exacta: "Aprendimos a utilizar la IA como un par de programación de nivel senior. Nos enseñó a implementar la arquitectura de Signals en Angular 22, a consumir Web APIs nativas (cámara y sonido) sin librerías pesadas y a redactar bitácoras de prompts (prompts.md) para estructurar software de forma transparente."', bold=True)

add_heading_2('Pregunta 6: "¿Qué viene para las Semanas 5 a 8?"')
add_body('Respuesta Exacta: "Iniciamos la fase comercial bajo el modelo \'Regalo + Comunidad\' (Semanas 5 y 6) con la meta individual de 2 usuarios reales por Doer. En la Semana 7 iteraremos según el feedback en NotebookLM y en la Semana 8 presentaremos las métricas finales en el Demo Day."', bold=True)

# 5. ANEXO PROMPTS.MD
add_heading_1('5. ANEXO: CONTENIDO PARA EL ARCHIVO prompts.md DEL REPOSITORIO')
add_body('Copia este texto y guárdalo en un archivo llamado prompts.md en la raíz de tu proyecto o en Notion:')
add_body(
'# Bitácora de Prompts — CheckUp+ (Sprint 16 / Semana 4)\n\n'
'## Prompt 1: Generación de Pase Verde con QR Dinámico y Temporizador\n'
'• Propósito: Implementar la lógica del Pase Verde con temporizador reactivo y cámara.\n'
'• Prompt: "Genera un componente standalone en Angular 22 (green-pass.ts) utilizando Signals para el manejo del estado. Debe incluir un temporizador en tiempo real de 15 minutos que se reinicie al expirar. Integra un escáner de cámara web con navigator.mediaDevices.getUserMedia, feedback de audio con AudioContext al validar, y modal de simulación de resultado sin exponer diagnósticos (RF4.2)."\n\n'
'## Prompt 2: Sistema de Agendamiento Express en < 3 Clics\n'
'• Propósito: Cumplir con el requerimiento RF2.2 de Reserva Express.\n'
'• Prompt: "Diseña el flujo de agendamiento express dentro de home.ts. Permite filtrar laboratorios por categoría (\'rutina\', \'cercanas\'), seleccionar fecha, horario y tipos de exámenes en un modal fluido con spinners de carga y confirmación en menos de 3 pasos."\n\n'
'## Prompt 3: Estilizado Glassmorphism & UI Confidencial\n'
'• Propósito: Aplicar la identidad visual definida en DESIGN.md.\n'
'• Prompt: "Aplica estilos CSS en app.css bajo principios de Glassmorphism, paleta de colores oscura segura para salud sexual y tarjetas de indicadores KPI legibles cumpliendo estándar WCAG AA."',
italic=True)

out_path = r'C:/Users/USER/Desktop/Guia_Presentacion_CheckUp_Sabado.docx'
doc.save(out_path)
print('SUCCESS_SAVED:', out_path)
